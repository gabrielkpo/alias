import docx
import pytest
from pypdf import PdfReader, PdfWriter

from alias_pii import readers
from alias_pii.masking import find_leaks, mask
from alias_pii.readers import UnsupportedFileError, read_document
from alias_pii.rules import RuleDetector

IBAN = "FR76 3000 6000 0112 3456 7890 189"


def test_plain_text_and_markdown(tmp_path):
    (tmp_path / "a.txt").write_text("Écrire à a.b@example.com", encoding="utf-8")
    (tmp_path / "b.md").write_bytes("﻿# Note\nTél. 01 99 00 12 34".encode())
    txt, md = read_document(tmp_path / "a.txt"), read_document(tmp_path / "b.md")
    assert (txt.format, txt.text, txt.warnings) == ("txt", "Écrire à a.b@example.com", [])
    assert md.text == "# Note\nTél. 01 99 00 12 34"  # byte order mark dropped


def test_legacy_encoding_is_decoded_and_reported(tmp_path):
    (tmp_path / "old.txt").write_bytes("Hélène écrit".encode("cp1252"))
    document = read_document(tmp_path / "old.txt")
    assert document.text == "Hélène écrit"
    assert document.warnings == ["not valid UTF-8, decoded as Windows-1252"]


def test_docx_keeps_headers_footers_and_tables(tmp_path):
    word = docx.Document()
    word.sections[0].header.paragraphs[0].text = "Dossier n° 2024/00421"
    word.sections[0].footer.paragraphs[0].text = "contact@example.com"
    word.add_paragraph("Bonjour,")
    table = word.add_table(rows=1, cols=3)
    table.cell(0, 0).text = "IBAN"
    table.cell(0, 1).merge(table.cell(0, 2)).text = IBAN
    word.add_paragraph("Cordialement")
    word.add_section()  # its header is linked to the first one: must not be read twice
    word.save(tmp_path / "letter.docx")

    document = read_document(tmp_path / "letter.docx")
    assert document.text.splitlines() == [
        "Dossier n° 2024/00421",
        "contact@example.com",
        "Bonjour,",
        f"IBAN | {IBAN}",
        "Cordialement",
    ]


def test_pdf_pages_are_kept_apart_and_empty_pages_reported(make_pdf):
    path = make_pdf([["Madame (cliente)", f"IBAN {IBAN}"], [], ["contact@example.com"]])
    document = read_document(path)
    pages = document.text.split("\f")
    assert len(pages) == 3
    assert pages[0].splitlines() == ["Madame (cliente)", f"IBAN {IBAN}"]
    assert pages[2].strip() == "contact@example.com"
    assert document.warnings == [
        "page 2: no extractable text (scanned page? OCR is not supported yet)"
    ]


def test_password_protected_pdf_is_refused(make_pdf, tmp_path):
    writer = PdfWriter(clone_from=PdfReader(make_pdf([["secret"]])))
    writer.encrypt("pass", algorithm="RC4-128")
    writer.write(tmp_path / "locked.pdf")
    with pytest.raises(UnsupportedFileError, match="password"):
        read_document(tmp_path / "locked.pdf")


def test_broken_pdf_is_refused(tmp_path):
    (tmp_path / "broken.pdf").write_bytes(b"not a pdf at all")
    with pytest.raises(UnsupportedFileError, match="unreadable"):
        read_document(tmp_path / "broken.pdf")


@pytest.mark.parametrize("name", ["scan.png", "notes.odt", "README"])
def test_unsupported_types_are_refused(tmp_path, name):
    (tmp_path / name).write_bytes(b"x")
    with pytest.raises(UnsupportedFileError, match="unsupported"):
        read_document(tmp_path / name)


def test_oversized_file_is_refused(tmp_path, monkeypatch):
    monkeypatch.setattr(readers, "MAX_BYTES", 10)
    (tmp_path / "big.txt").write_text("x" * 11)
    with pytest.raises(UnsupportedFileError, match="limit"):
        read_document(tmp_path / "big.txt")


def test_repr_does_not_show_the_text(tmp_path):
    (tmp_path / "a.txt").write_text("a.b@example.com")
    assert "a.b@example.com" not in repr(read_document(tmp_path / "a.txt"))


def test_read_then_mask_leaves_nothing_in_clear(tmp_path):
    word = docx.Document()
    word.sections[0].footer.paragraphs[0].text = "a.b@example.com"
    word.add_table(rows=1, cols=1).cell(0, 0).text = IBAN
    word.save(tmp_path / "w.docx")
    result = mask(read_document(tmp_path / "w.docx").text, RuleDetector())
    assert result.text == "[EMAIL_1]\n[IBAN_1]"
    assert find_leaks(result.text, result.entities) == []


@pytest.mark.parametrize("encoding", ["utf-16", "utf-32"])
def test_utf16_and_utf32_with_byte_order_mark(tmp_path, encoding):
    (tmp_path / "unicode.txt").write_bytes("Écrire à a.b@example.com".encode(encoding))
    document = read_document(tmp_path / "unicode.txt")
    assert (document.text, document.warnings) == ("Écrire à a.b@example.com", [])


def test_utf16_without_byte_order_mark_is_not_read_as_windows_1252(tmp_path):
    (tmp_path / "nobom.txt").write_bytes("mail a.b@example.com".encode("utf-16-le"))
    document = read_document(tmp_path / "nobom.txt")
    assert document.text == "mail a.b@example.com"
    assert document.warnings == ["no byte order mark, decoded as utf-16-le"]
    assert mask(document.text, RuleDetector()).text == "mail [EMAIL_1]"


def _append_xml(word, xml):
    from docx.oxml import parse_xml

    word.element.body.insert(len(word.element.body) - 1, parse_xml(xml))


W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
MC = 'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'


def test_docx_text_boxes_content_controls_and_comments(tmp_path):
    word = docx.Document()
    word.add_paragraph("Bonjour,")
    # A text box is stored twice (Choice and legacy Fallback): it must be read once.
    box = "<w:txbxContent><w:p><w:r><w:t>{}</w:t></w:r></w:p></w:txbxContent>"
    _append_xml(
        word,
        f"<w:p {W} {MC}><w:r><mc:AlternateContent>"
        f"<mc:Choice Requires='wps'>{box.format('12 rue des Lilas')}</mc:Choice>"
        f"<mc:Fallback>{box.format('12 rue des Lilas')}</mc:Fallback>"
        "</mc:AlternateContent></w:r></w:p>",
    )
    _append_xml(
        word,
        f"<w:sdt {W}><w:sdtContent><w:p><w:r><w:t>Client : a.b@example.com</w:t></w:r></w:p>"
        "</w:sdtContent></w:sdt>",
    )
    signed = word.add_paragraph("Signé")
    word.add_comment(signed.runs, text="Relu par c.d@example.com", author="X")
    word.save(tmp_path / "w.docx")

    lines = read_document(tmp_path / "w.docx").text.splitlines()
    assert lines.count("12 rue des Lilas") == 1
    assert "Client : a.b@example.com" in lines
    assert "Relu par c.d@example.com" in lines


def test_docx_equal_adjacent_cells_are_kept(tmp_path):
    word = docx.Document()
    row = word.add_table(rows=1, cols=3).rows[0]
    for cell, value in zip(row.cells, ["Oui", "Oui", "Non"], strict=True):
        cell.text = value
    word.save(tmp_path / "t.docx")
    assert read_document(tmp_path / "t.docx").text == "Oui | Oui | Non"


def test_broken_docx_is_refused(tmp_path):
    (tmp_path / "old.docx").write_bytes(b"\xd0\xcf\x11\xe0 legacy .doc renamed")
    with pytest.raises(UnsupportedFileError, match="not a readable .docx"):
        read_document(tmp_path / "old.docx")


def test_pdf_form_fields_and_annotations_are_read(make_pdf):
    path = make_pdf(
        [["Nom :", "IBAN :"]],
        fields={"nom": "Hélène Martin", "iban": IBAN},
        note="Rappeler au 01 99 00 12 34",
    )
    extra = read_document(path).text.split("\f")[-1].splitlines()
    assert extra == ["nom: Hélène Martin", f"iban: {IBAN}", "Rappeler au 01 99 00 12 34"]


def test_one_broken_pdf_page_does_not_hide_the_others(make_pdf, monkeypatch):
    from pypdf import PageObject

    original = PageObject.extract_text

    def flaky(self, *args, **kwargs):
        if "page two" in original(self, *args, **kwargs):
            raise KeyError("/Font")
        return original(self, *args, **kwargs)

    monkeypatch.setattr(PageObject, "extract_text", flaky)
    document = read_document(make_pdf([["page one"], ["page two"], ["page three"]]))
    assert [p.strip() for p in document.text.split("\f")] == ["page one", "", "page three"]
    assert document.warnings == ["page 2: text could not be extracted (KeyError)"]


def test_binary_file_with_a_text_suffix_is_refused(tmp_path):
    (tmp_path / "image.txt").write_bytes(b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x01")
    with pytest.raises(UnsupportedFileError, match="NUL bytes"):
        read_document(tmp_path / "image.txt")


def test_big_endian_utf16_without_byte_order_mark(tmp_path):
    (tmp_path / "be.txt").write_bytes("Écrire à a.b@example.com".encode("utf-16-be"))
    document = read_document(tmp_path / "be.txt")
    assert document.text == "Écrire à a.b@example.com"
    assert document.warnings == ["no byte order mark, decoded as utf-16-be"]


@pytest.mark.parametrize(
    "raw",
    [
        b"Contact: a.b@example.com, tel 01 99 00 12 34\x00",  # UTF-8 with a stray NUL
        b"Contact: a.b@example.com\x00\x00",  # padding
    ],
)
def test_text_with_stray_nul_bytes_is_refused_not_mis_decoded(tmp_path, raw):
    """Decoded as UTF-16 it would become CJK-looking text: unmasked yet fully recoverable."""
    (tmp_path / "odd.txt").write_bytes(raw)
    with pytest.raises(UnsupportedFileError, match="NUL bytes"):
        read_document(tmp_path / "odd.txt")
